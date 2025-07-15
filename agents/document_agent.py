import os
from PyPDF2 import PdfReader
import docx
import openpyxl

def run_document_parser(uploaded_file):
    ext = os.path.splitext(uploaded_file.name)[-1].lower()

    try:
        if ext == ".pdf":
            reader = PdfReader(uploaded_file)
            text = "\n".join([page.extract_text() or "" for page in reader.pages])
            return text.strip()

        elif ext == ".docx":
            doc = docx.Document(uploaded_file)
            return "\n".join([p.text for p in doc.paragraphs])

        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(uploaded_file, data_only=True)
            sheet = wb.active
            data = []
            for row in sheet.iter_rows(values_only=True):
                data.append("\t".join([str(cell) if cell is not None else "" for cell in row]))
            return "\n".join(data)

        else:
            return "Unsupported file type."
    except Exception as e:
        return f"❌ Error parsing document: {e}"
